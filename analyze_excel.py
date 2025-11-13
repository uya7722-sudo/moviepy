import openpyxl
import pandas as pd
import json

# Load the Excel file
file_path = '/vercel/sandbox/uploads/Aircraft_Jobcard_Database.xlsm'
wb = openpyxl.load_workbook(file_path, data_only=True)

print("=== SHEET NAMES ===")
print(wb.sheetnames)
print()

# Analyze each sheet
analysis = {}
for sheet_name in wb.sheetnames:
    print(f"=== ANALYZING SHEET: {sheet_name} ===")
    ws = wb[sheet_name]
    
    # Get dimensions
    max_row = ws.max_row
    max_col = ws.max_column
    print(f"Dimensions: {max_row} rows x {max_col} columns")
    
    # Read data using pandas
    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name, engine='openpyxl')
        print(f"Columns: {list(df.columns)}")
        print(f"Data shape: {df.shape}")
        print(f"First few rows:")
        print(df.head(10))
        print()
        
        # Store analysis
        analysis[sheet_name] = {
            'columns': list(df.columns),
            'shape': df.shape,
            'dtypes': df.dtypes.astype(str).to_dict(),
            'sample_data': df.head(5).to_dict('records')
        }
    except Exception as e:
        print(f"Error reading sheet: {e}")
        print()

# Save analysis to JSON
with open('/vercel/sandbox/excel_analysis.json', 'w') as f:
    json.dump(analysis, f, indent=2, default=str)

print("Analysis saved to excel_analysis.json")
